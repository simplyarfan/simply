"""Turn the corpus folder into plain text, one record per file.

Assume the folder is hostile. The graded one has an empty dir, a chmod 000
file, an encrypted PDF and a junk binary, and there may be more surprises.
Every file is handled on its own: if it blows up we note why and move on.
"""

from __future__ import annotations

import csv
import io
import os
import re
import signal
import threading
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path

TEXT_EXT = {
    ".txt", ".log", ".md", ".rst", ".py", ".json", ".yaml", ".yml", ".toml",
    ".ini", ".cfg", ".conf", ".sh", ".sql", ".xml", ".html", ".htm", ".js",
    ".ts", ".c", ".h", ".cpp", ".hpp", ".java", ".go", ".rs",
}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
TABLE_EXT = {".csv", ".tsv"}

# a page with less text than this is probably a scan, so we OCR its images
SCANNED_PAGE_CHARS = 25
# figures on normal pages get read too (time allowing) if they're big enough
# to hold readable text, a few per document at most
FIGURE_MIN_SIDE = 160
FIGURES_PER_PDF = 4
PER_FILE_SECONDS = 20

# words that mean "don't trust this document"
STALE_RE = re.compile(
    r"\b(withdrawn|superseded|obsolete|deprecated|do not use)\b", re.I
)


@dataclass
class Doc:
    rel: str                      # corpus-relative path, forward slashes
    kind: str                     # pdf, docx, xlsx, table, text, image
    text: str = ""
    flags: list[str] = field(default_factory=list)
    # images still waiting for the vision model: (label, PIL image, priority).
    # Priority 0 is a picture file, 1 a scanned PDF page, 2 a figure on a page
    # that already has text. If startup time runs short, 2 is what we drop.
    pending_images: list = field(default_factory=list)


@dataclass
class Skipped:
    rel: str
    reason: str


class FileTimeout(Exception):
    pass


@contextmanager
def time_limit(seconds: int):
    # signal.alarm only works in the main thread. The server parses in the
    # main thread so we get the real thing there; anywhere else it's a no-op.
    if threading.current_thread() is not threading.main_thread():
        yield
        return

    def _boom(signum, frame):
        raise FileTimeout(f"took longer than {seconds}s")

    old = signal.signal(signal.SIGALRM, _boom)
    signal.alarm(seconds)
    try:
        yield
    finally:
        signal.alarm(0)
        signal.signal(signal.SIGALRM, old)


def walk(root: Path) -> list[Path]:
    """Every regular file under root, sorted so runs are repeatable."""
    found = []

    def _complain(err):  # unreadable directory: skip it, keep walking
        print(f"[corpus] can't list {err.filename}: {err}", flush=True)

    for dirpath, dirnames, filenames in os.walk(root, onerror=_complain, followlinks=False):
        dirnames.sort()
        for name in sorted(filenames):
            p = Path(dirpath) / name
            try:
                if p.is_file() and not p.is_symlink():
                    found.append(p)
            except OSError:
                pass
    return found


def load_corpus(root: Path) -> tuple[list[Doc], list[Skipped]]:
    docs, skipped = [], []
    for path in walk(root):
        rel = path.relative_to(root).as_posix()
        try:
            with time_limit(PER_FILE_SECONDS):
                doc = read_one(path, rel)
        except FileTimeout as e:
            skipped.append(Skipped(rel, str(e)))
            continue
        except Exception as e:  # noqa: BLE001 - genuinely anything can come out of a bad file
            skipped.append(Skipped(rel, f"{type(e).__name__}: {e}"[:200]))
            continue
        if isinstance(doc, Skipped):
            skipped.append(doc)
        else:
            doc.flags = stale_flags(rel, doc.text)
            docs.append(doc)
    return docs, skipped


def read_one(path: Path, rel: str) -> Doc | Skipped:
    ext = path.suffix.lower()

    # Open it once up front. A chmod 000 file fails right here with
    # PermissionError, before any parser gets a chance to do something odd.
    with open(path, "rb") as fh:
        head = fh.read(8)

    if ext == ".pdf":
        return read_pdf(path, rel)
    if ext == ".docx":
        return read_docx(path, rel)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path, rel)
    if ext in TABLE_EXT:
        return read_table(path, rel, "\t" if ext == ".tsv" else None)
    if ext in IMAGE_EXT:
        return read_image(path, rel)
    if ext in TEXT_EXT:
        return read_text(path, rel)
    if ext == "" and not looks_binary(path):
        return read_text(path, rel)
    return Skipped(rel, f"no parser for {ext or 'extensionless'} file ({head[:4]!r})")


def looks_binary(path: Path) -> bool:
    with open(path, "rb") as fh:
        chunk = fh.read(4096)
    return b"\x00" in chunk


def stale_flags(rel: str, text: str) -> list[str]:
    flags = []
    # Only the filename and the title lines count. A current datasheet often
    # says "revision 1 is withdrawn" further down, and flagging THAT one as
    # stale would steer the model straight to the wrong file.
    top = [ln for ln in text.splitlines() if ln.strip() and not ln.startswith("[page")][:4]
    hit = STALE_RE.search(Path(rel).name.replace("_", " ")) or STALE_RE.search("\n".join(top))
    if hit:
        flags.append(f"marked {hit.group(1).upper()}")
    return flags


# ---- individual formats -----------------------------------------------------

def read_text(path: Path, rel: str) -> Doc | Skipped:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        return Skipped(rel, "binary content")
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return Doc(rel, "text", text.replace("\r\n", "\n"))


def read_pdf(path: Path, rel: str) -> Doc | Skipped:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        # Even if an empty password would open it, the rules say a locked file
        # is not a source. Anything we "found" in there would score as a guess.
        return Skipped(rel, "encrypted pdf")

    parts, pending, figures = [], [], 0
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as e:  # noqa: BLE001
            text = ""
            print(f"[corpus] {rel} page {i}: text extraction failed: {e}", flush=True)
        parts.append(f"[page {i}]\n{text.strip()}")
        scanned = len(text.strip()) < SCANNED_PAGE_CHARS
        try:
            for j, img in enumerate(page.images):
                im = img.image
                if scanned:
                    # probably a scanned page: its images ARE the page
                    pending.append((f"page {i} image {j + 1}", im.convert("RGB"), 1))
                elif figures < FIGURES_PER_PDF and min(im.size) >= FIGURE_MIN_SIDE:
                    # a diagram on a normal page can carry text nothing else has
                    pending.append((f"page {i} figure {j + 1}", im.convert("RGB"), 2))
                    figures += 1
        except Exception as e:  # noqa: BLE001
            print(f"[corpus] {rel} page {i}: couldn't pull images: {e}", flush=True)

    return Doc(rel, "pdf", "\n\n".join(parts), pending_images=pending)


def read_docx(path: Path, rel: str) -> Doc:
    import docx
    from docx.table import Table

    d = docx.Document(str(path))
    lines = []

    # body in reading order, so a table stays next to the heading above it
    for block in d.iter_inner_content():
        if isinstance(block, Table):
            lines.extend(table_lines(block))
        else:
            t = block.text.strip()
            if t:
                lines.append(t)

    for section in d.sections:
        for part in (section.header, section.footer):
            try:
                extra = [p.text.strip() for p in part.paragraphs if p.text.strip()]
            except Exception:  # noqa: BLE001
                extra = []
            if extra:
                lines.append("[header/footer] " + " / ".join(extra))

    # Text boxes and tables nested inside table cells aren't in either list
    # above. Sweep every paragraph in the XML and keep whatever we missed.
    from docx.oxml.ns import qn
    seen = " ".join(" ".join(lines).split())
    for p in d.element.body.iter(qn("w:p")):
        t = " ".join("".join(x.text or "" for x in p.iter(qn("w:t"))).split())
        if t and t not in seen:
            lines.append(t)
            seen += " " + t

    return Doc(rel, "docx", "\n".join(lines))


def table_lines(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        for c in row.cells:
            t = " ".join(c.text.split())
            # merged cells come back once per grid column, drop the repeats
            if not cells or cells[-1] != t:
                cells.append(t)
        if any(cells):
            rows.append(cells)
    return labelled_rows(rows, prefix="[table]")


def labelled_rows(rows: list[list[str]], prefix: str) -> list[str]:
    """Write rows as 'Header: value | Header: value'.

    A bare row like '4.3.2 | closed | S2' means nothing to the model once it's
    out of the grid, so every value carries its column name with it.
    """
    if not rows:
        return []
    header = rows[0]
    out = [f"{prefix} columns: " + " | ".join(header)]
    for r in rows[1:]:
        if len(r) == len(header) and any(header):
            out.append(prefix + " " + " | ".join(f"{h}: {v}" for h, v in zip(header, r) if v != ""))
        else:
            out.append(prefix + " " + " | ".join(r))
    return out


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        s = v.isoformat()
        return s[:10] if s.endswith("T00:00:00") else s
    return str(v).strip()


def read_xlsx(path: Path, rel: str) -> Doc:
    import openpyxl

    values = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    # Files saved by scripts often have formulas but no cached results, so
    # data_only gives None. Fall back to showing the formula itself.
    formulas = openpyxl.load_workbook(str(path), read_only=True, data_only=False)

    lines = []
    for ws in values.worksheets:
        fws = formulas[ws.title]
        rows = []
        for vrow, frow in zip(ws.iter_rows(values_only=True), fws.iter_rows(values_only=True)):
            cells = [cell_str(v if v is not None else f) for v, f in zip(vrow, frow)]
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                rows.append(cells)
        if rows:
            lines.append(f"[sheet: {ws.title}]")
            lines.extend(labelled_rows(rows, prefix=f"[{ws.title}]"))
    values.close()
    formulas.close()
    return Doc(rel, "xlsx", "\n".join(lines))


def read_table(path: Path, rel: str, delimiter: str | None) -> Doc:
    raw = path.read_bytes().decode("utf-8-sig", errors="replace")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(raw), delimiter=delimiter)]
    rows = [r for r in rows if any(r)]
    return Doc(rel, "table", "\n".join(labelled_rows(rows, prefix="[row]")))


def read_image(path: Path, rel: str) -> Doc:
    from PIL import Image, ImageOps

    Image.MAX_IMAGE_PIXELS = 400_000_000  # a big scan is not a decompression bomb

    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)  # phone photos are often stored sideways
        im = im.convert("RGB")
        im.load()
    return Doc(rel, "image", "", pending_images=[("image", im, 0)])
